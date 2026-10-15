import os
import io
import time
import tempfile
from dotenv import load_dotenv
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_google_genai import GoogleGenerativeAIEmbeddings
import chromadb

from pypdf import PdfReader
import docx
from pptx import Presentation
import openpyxl

def extract_documents_from_file(temp_path, ext, base_meta):
    """파일 확장자별로 텍스트를 추출합니다. 청크 크기를 대폭 키워 API 호출을 최소화합니다."""
    docs = []
    # [최적화 1] 1000 -> 2500으로 키워 조각 개수를 절반 이하로 줄임 (일일 한도 절약)
    text_splitter = RecursiveCharacterTextSplitter(chunk_size=2500, chunk_overlap=200)

    if ext == 'pdf':
        reader = PdfReader(temp_path)
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text and text.strip():
                meta = base_meta.copy()
                meta["location"] = f"{i + 1}페이지"
                for chunk in text_splitter.split_text(text):
                    docs.append(Document(page_content=chunk, metadata=meta))

    elif ext == 'docx':
        doc = docx.Document(temp_path)
        full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        if full_text:
            meta = base_meta.copy()
            meta["location"] = "본문"
            for chunk in text_splitter.split_text(full_text):
                docs.append(Document(page_content=chunk, metadata=meta))

    elif ext == 'pptx':
        prs = Presentation(temp_path)
        for i, slide in enumerate(prs.slides):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            slide_texts.append(paragraph.text.strip())
            full_slide_text = "\n".join(slide_texts)
            if full_slide_text:
                meta = base_meta.copy()
                meta["location"] = f"슬라이드 {i + 1}장"
                for chunk in text_splitter.split_text(full_slide_text):
                    docs.append(Document(page_content=chunk, metadata=meta))

    elif ext == 'xlsx':
        wb = openpyxl.load_workbook(temp_path, data_only=True)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                continue
            headers = [str(h).strip() if h is not None else f"열{idx+1}" for idx, h in enumerate(rows[0])]
            
            batch_rows_text = []
            start_row_num = 2
            
            # [최적화 2] 엑셀은 10행 -> 20행 단위로 묶어 조각 개수를 더욱 감축
            for r_idx, row in enumerate(rows[1:], start=2):
                row_items = []
                for h, val in zip(headers, row):
                    if val is not None and str(val).strip() != "":
                        row_items.append(f"{h}: {str(val).strip()}")
                if row_items:
                    batch_rows_text.append(f"({r_idx}행) " + " | ".join(row_items))
                
                if len(batch_rows_text) >= 20 or r_idx == len(rows):
                    if batch_rows_text:
                        combined_text = f"[{sheet} 시트 {start_row_num}~{r_idx}행]\n" + "\n".join(batch_rows_text)
                        meta = base_meta.copy()
                        meta["location"] = f"{sheet} 시트 {start_row_num}~{r_idx}행"
                        docs.append(Document(page_content=combined_text, metadata=meta))
                        batch_rows_text = []
                        start_row_num = r_idx + 1

    return docs


def embed_with_rate_limit(embeddings, texts, batch_size=100):
    """무한 루프(Hang) 방지 및 API 효율 극대화를 위한 임베딩 함수"""
    all_vectors = []
    total = len(texts)
    max_retries = 3  # [안전장치] 최대 3회 재시도 후 에러를 내뿜고 강제 종료
    
    # [최적화 3] 한 번에 최대치인 100개씩 묶어서 구글 서버로 전송
    for i in range(0, total, batch_size):
        batch = texts[i : i + batch_size]
        retry_count = 0
        
        while True:
            try:
                vectors = embeddings.embed_documents(batch)
                all_vectors.extend(vectors)
                
                # [안심 기능] 터미널에 실시간 진행률 표시
                print(f"     ... 임베딩 진행률: {min(i + batch_size, total)} / {total} 조각 완료")
                
                if i + batch_size < total:
                    time.sleep(2) 
                break
            except Exception as e:
                err_msg = str(e)
                if "429" in err_msg or "RESOURCE_EXHAUSTED" in err_msg:
                    retry_count += 1
                    
                    if retry_count > max_retries:
                        print("\n[치명적 오류] 일일 무료 할당량(Daily Quota)이 고갈되었거나 네트워크가 차단되었습니다.")
                        print(" -> 무한 루프를 방지하기 위해 작업을 강제 종료합니다. (수 시간 뒤 또는 내일 다시 시도하세요)")
                        raise e  # 무한루프를 끊고 강제 에러 발생
                        
                    wait_sec = 60  # 1분 대기
                    print(f"   -> [API 할당량 초과] {wait_sec}초 대기 후 재시도... (재시도 {retry_count}/{max_retries}회)")
                    time.sleep(wait_sec)
                else:
                    raise e
                    
    return all_vectors


def get_all_files_recursively(drive_service, folder_id):
    all_files = []
    folders_to_scan = [folder_id]

    while folders_to_scan:
        current_folder = folders_to_scan.pop(0)
        query = f"'{current_folder}' in parents and trashed = false"
        results = drive_service.files().list(
            q=query,
            pageSize=1000,
            fields="files(id, name, mimeType, fileExtension, webViewLink, modifiedTime)"
        ).execute()

        for item in results.get('files', []):
            if item.get('mimeType') == 'application/vnd.google-apps.folder':
                folders_to_scan.append(item.get('id'))
            else:
                all_files.append(item)

    return all_files


def ingest_manuals_to_chroma():
    print("=== [Step 1-B: 매뉴얼 및 BIM 엑셀 문서 Vector DB 이중 연동 동기화 시작] ===\n")

    load_dotenv()
    key_path = os.getenv("GOOGLE_APPLICATION_CREDENTIALS")
    manual_folder_id = os.getenv("MANUAL_FOLDER_ID")
    bim_folder_id = os.getenv("BIM_FILE_FOLDER_ID")  # [추가] BIM 파일 폴더 ID 로드
    chroma_path = os.getenv("CHROMA_DB_PATH")
    gemini_api_key = os.getenv("GOOGLE_API_KEY")

    chroma_client = chromadb.PersistentClient(path=chroma_path)
    collection = chroma_client.get_or_create_collection(name="bim_manuals")

    existing_data = collection.get(include=["metadatas"])
    existing_files = {}
    for meta in existing_data.get("metadatas", []):
        if meta and "file_id" in meta:
            existing_files[meta["file_id"]] = meta.get("modified_time", "")

    print(f"1. 현재 Vector DB에 등록된 문서 파일 수: {len(existing_files)}개")

    drive_scope = "https://" + "www.googleapis.com/auth/drive.readonly"
    creds = Credentials.from_service_account_file(key_path, scopes=[drive_scope])
    drive_service = build('drive', 'v3', credentials=creds)

    # 2. [핵심 수정] 매뉴얼 폴더와 BIM 폴더 양쪽을 모두 스캔하여 파일 목록 병합
    print("2. 구글 드라이브 스캔 중 (매뉴얼 폴더 및 BIM 파일 폴더 동시 탐색)...")
    manual_files = get_all_files_recursively(drive_service, manual_folder_id)
    bim_files = get_all_files_recursively(drive_service, bim_folder_id)
    
    # 두 폴더에서 스캔된 파일들을 하나의 딕셔너리로 합쳐 중복 제거
    unique_drive_files = {f.get('id'): f for f in manual_files + bim_files}
    drive_files = list(unique_drive_files.values())
    
    print(f" -> 총 {len(drive_files)}개 파일 스캔 완료 (지원 포맷 필터링 전)\n")

    drive_file_ids = set(unique_drive_files.keys())
    deleted_ids = set(existing_files.keys()) - drive_file_ids
    deleted_count = 0

    if deleted_ids:
        print("3. [삭제 감지] 구글 드라이브에서 제거된 파일 데이터를 Vector DB에서 지웁니다.")
        for del_id in deleted_ids:
            collection.delete(where={"file_id": del_id})
            deleted_count += 1
        print(f" - 총 {deleted_count}개 문서 데이터 정리 완료\n")

    embeddings = GoogleGenerativeAIEmbeddings(
        model="models/gemini-embedding-001",
        google_api_key=gemini_api_key
    )

    supported_extensions = {'pdf', 'docx', 'pptx', 'xlsx'}
    new_count = 0
    updated_count = 0

    for file in drive_files:
        file_id = file.get('id')
        file_name = file.get('name', '')
        mime_type = file.get('mimeType', '')
        drive_mod_time = file.get('modifiedTime', '')

        ext = file.get('fileExtension', '').lower()
        if not ext and '.' in file_name:
            ext = file_name.rsplit('.', 1)[-1].lower()

        export_mime = None
        if mime_type == 'application/vnd.google-apps.spreadsheet':
            ext = 'xlsx'
            export_mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        elif mime_type == 'application/vnd.google-apps.presentation':
            ext = 'pptx'
            export_mime = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
        elif mime_type == 'application/vnd.google-apps.document':
            ext = 'docx'
            export_mime = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'

        # .rfa, .rvt 등 지원하지 않는 파일은 여기서 모두 안전하게 걸러짐
        if ext not in supported_extensions:
            continue

        is_update = False
        if file_id in existing_files:
            if existing_files[file_id] == drive_mod_time:
                print(f" - [유지] '{file_name}' (변경사항 없음)")
                continue
            else:
                print(f" - [변경 감지] '{file_name}' 재학습 중...")
                collection.delete(where={"file_id": file_id})
                is_update = True
        else:
            print(f" - [신규 처리] '{file_name}' 텍스트 추출 중...")

        try:
            if export_mime:
                request = drive_service.files().export_media(fileId=file_id, mimeType=export_mime)
            else:
                request = drive_service.files().get_media(fileId=file_id)
                
            fh = io.BytesIO()
            downloader = MediaIoBaseDownload(fh, request)
            done = False
            while not done:
                status, done = downloader.next_chunk()

            with tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}") as temp_file:
                temp_file.write(fh.getvalue())
                temp_path = temp_file.name

            base_meta = {
                "file_id": file_id,
                "file_name": file_name,
                "drive_link": file.get('webViewLink', ''),
                "modified_time": drive_mod_time
            }
            
            docs = extract_documents_from_file(temp_path, ext, base_meta)

            if docs:
                texts = [doc.page_content for doc in docs]
                metadatas = [doc.metadata for doc in docs]
                ids = [f"{file_id}_chunk_{i}" for i in range(len(docs))]

                print(f"   -> 추출된 덩어리 수: {len(docs)}개 (서버 전송 시작...)")
                
                vectors = embed_with_rate_limit(embeddings, texts, batch_size=100)
                
                collection.add(
                    ids=ids,
                    embeddings=vectors,
                    documents=texts,
                    metadatas=metadatas
                )
                if is_update:
                    updated_count += 1
                else:
                    new_count += 1
                print(f"   -> 완료: '{file_name}' (DB 적재 성공)\n")
            else:
                print(f"   -> [경고] '{file_name}'에서 읽을 수 있는 텍스트가 없습니다.\n")

        except Exception as e:
            print(f"   -> [오류 발생] '{file_name}' 처리 실패: {e}\n")

        finally:
            if 'temp_path' in locals() and os.path.exists(temp_path):
                os.remove(temp_path)

    print(f"\n4. 작업 요약: 신규 {new_count}개 / 수정 {updated_count}개 / 삭제 {deleted_count}개")
    print("=== [Vector DB 완전 동기화 종료] ===")

if __name__ == "__main__":
    ingest_manuals_to_chroma()